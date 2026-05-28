"""
Simplified CPI-to-AWS Connectivity Requirements Questionnaire.
Focus: What we need FROM the client to establish connectivity.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def create_styles():
    """Create reusable styles"""
    return {
        "header": {
            "font": Font(bold=True, color="FFFFFF", size=11),
            "fill": PatternFill(
                start_color="1F4E79", end_color="1F4E79", fill_type="solid"
            ),
            "alignment": Alignment(
                horizontal="center", vertical="center", wrap_text=True
            ),
            "border": Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            ),
        },
        "section": {
            "font": Font(bold=True, color="FFFFFF", size=12),
            "fill": PatternFill(
                start_color="2E75B6", end_color="2E75B6", fill_type="solid"
            ),
            "alignment": Alignment(horizontal="left", vertical="center"),
            "border": Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            ),
        },
        "data": {
            "font": Font(size=10),
            "alignment": Alignment(
                horizontal="left", vertical="center", wrap_text=True
            ),
            "border": Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            ),
        },
    }


def apply_style(cell, style_dict):
    for attr, value in style_dict.items():
        setattr(cell, attr, value)


def add_section(ws, row, title, styles, cols=3):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row=row, column=1, value=title)
    apply_style(cell, styles["section"])
    ws.row_dimensions[row].height = 25
    return row + 1


def add_header(ws, row, headers, styles):
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        apply_style(cell, styles["header"])
    ws.row_dimensions[row].height = 25
    return row + 1


def add_row(ws, row, data, styles):
    for col, value in enumerate(data, 1):
        cell = ws.cell(row=row, column=col, value=value)
        apply_style(cell, styles["data"])
    ws.row_dimensions[row].height = 35
    return row + 1


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "CPI Connectivity Requirements"
    styles = create_styles()
    row = 1

    # Title
    ws.merge_cells("A1:C1")
    title = ws.cell(
        row=1, column=1, value="RRPS CPI-to-Integrum AWS Connectivity Requirements"
    )
    title.font = Font(bold=True, size=16, color="1F4E79")
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35
    row = 3

    # Instructions
    ws.merge_cells("A3:C3")
    instr = ws.cell(
        row=3,
        column=1,
        value="Please complete the 'RRPS Response' column for each item below.",
    )
    instr.font = Font(italic=True, size=10)
    row = 5

    # ============ SECTION 1: Network ============
    row = add_section(ws, row, "1. Network Connectivity", styles)
    row = add_header(ws, row, ["#", "Requirement", "RRPS Response"], styles)

    network_items = [
        ("1.1", "CPI outbound IP range/CIDR (for our firewall whitelist)"),
        (
            "1.2",
            "Preferred connectivity method: Public Internet / VPN / AWS PrivateLink",
        ),
        ("1.3", "Any specific ports required beyond 443/HTTPS?"),
        ("1.4", "Network latency requirements (e.g., <100ms, <500ms)"),
    ]
    for item in network_items:
        row = add_row(ws, row, (item[0], item[1], ""), styles)

    row += 1

    # ============ SECTION 2: Authentication ============
    row = add_section(ws, row, "2. Authentication", styles)
    row = add_header(ws, row, ["#", "Requirement", "RRPS Response"], styles)

    auth_items = [
        (
            "2.1",
            "Preferred auth method: OAuth 2.0 / API Key / mTLS / Client Certificate",
        ),
        ("2.2", "If OAuth 2.0: Authorization Server URL"),
        ("2.3", "If OAuth 2.0: Required scopes for CPI"),
        ("2.4", "Token/credential rotation policy (e.g., every 90 days)"),
        ("2.5", "Will RRPS provide a dedicated service account?"),
    ]
    for item in auth_items:
        row = add_row(ws, row, (item[0], item[1], ""), styles)

    row += 1

    # ============ SECTION 3: TLS / Certificates ============
    row = add_section(ws, row, "3. TLS & Certificates", styles)
    row = add_header(ws, row, ["#", "Requirement", "RRPS Response"], styles)

    tls_items = [
        ("3.1", "Minimum TLS version required (We support TLS 1.3)"),
        ("3.2", "Is mutual TLS (mTLS) required?"),
        ("3.3", "Certificate authority preference: Public CA / Private CA"),
        ("3.4", "Certificate rotation frequency requirement"),
    ]
    for item in tls_items:
        row = add_row(ws, row, (item[0], item[1], ""), styles)

    row += 1

    # ============ SECTION 4: SAP CPI Details ============
    row = add_section(ws, row, "4. SAP CPI Configuration", styles)
    row = add_header(ws, row, ["#", "Requirement", "RRPS Response"], styles)

    cpi_items = [
        ("4.1", "CPI tenant/instance URL for integration"),
        ("4.2", "CPI environment: DEV / QA / PROD URLs"),
        ("4.3", "API rate limits from CPI (requests/minute)"),
        ("4.4", "Max payload size allowed"),
        ("4.5", "Content type: JSON / XML / Both"),
        ("4.6", "Retry policy for failed requests"),
    ]
    for item in cpi_items:
        row = add_row(ws, row, (item[0], item[1], ""), styles)

    row += 1

    # ============ SECTION 5: Data & Compliance ============
    row = add_section(ws, row, "5. Data & Compliance", styles)
    row = add_header(ws, row, ["#", "Requirement", "RRPS Response"], styles)

    data_items = [
        ("5.1", "Data residency requirement (e.g., APAC only, Singapore)"),
        ("5.2", "Any LLM provider restrictions? (OpenAI / Anthropic / Azure OpenAI)"),
        ("5.3", "Audit log retention requirement (years)"),
        ("5.4", "Security incident notification timeline (e.g., 24hrs)"),
    ]
    for item in data_items:
        row = add_row(ws, row, (item[0], item[1], ""), styles)

    row += 1

    # ============ SECTION 6: Contacts ============
    row = add_section(ws, row, "6. Key Contacts", styles)
    row = add_header(ws, row, ["#", "Role", "Name / Email"], styles)

    contact_items = [
        ("6.1", "CPI Technical Contact"),
        ("6.2", "Security/IT Contact"),
        ("6.3", "Project Sponsor"),
    ]
    for item in contact_items:
        row = add_row(ws, row, (item[0], item[1], ""), styles)

    # Column widths
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 45

    # Save
    output_path = "/Users/ssfoo/Documents/GitHub/rr_lead_to_cash/src/lead_to_cash/docs/reference_proposal/RRPS_CPI_Connectivity_Requirements.xlsx"
    wb.save(output_path)
    print(f"Created: {output_path}")
    return output_path


if __name__ == "__main__":
    main()
