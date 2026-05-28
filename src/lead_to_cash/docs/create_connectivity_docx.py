"""
Generate RRPS Connectivity Documentation in DOCX format.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import RGBColor


def set_cell_shading(cell, color):
    """Set cell background color"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc, headers, rows, header_color="1F4E79"):
    """Add a formatted table"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    # Header row
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        for para in header_cells[i].paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(header_cells[i], header_color)

    # Data rows
    for row_data in rows:
        row = table.add_row()
        for i, cell_data in enumerate(row_data):
            row.cells[i].text = str(cell_data)

    return table


def create_document():
    doc = Document()

    # Title
    title = doc.add_heading("Integrum API Connectivity Documentation", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata
    meta = doc.add_paragraph()
    meta.add_run("Version: ").bold = True
    meta.add_run("1.0 | ")
    meta.add_run("Date: ").bold = True
    meta.add_run("December 2025 | ")
    meta.add_run("Classification: ").bold = True
    meta.add_run("Confidential")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Section 1: Endpoint Details
    doc.add_heading("1. Endpoint Details", level=1)
    add_table(
        doc,
        ["Parameter", "Value"],
        [
            ["Base URL", "https://rr.kailash.ai"],
            ["Protocol", "HTTPS"],
            ["TLS Version", "1.3"],
            ["Port", "443"],
            ["Certificate", "Let's Encrypt (auto-renewal enabled)"],
        ],
    )

    doc.add_paragraph()

    # Section 2: Authentication
    doc.add_heading("2. Authentication", level=1)

    doc.add_heading("Testing (Current)", level=2)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Type", "HTTP Basic Authentication"],
            ["Username", "rrps_cpi_test"],
            ["Password", "Kailash2024CPI"],
        ],
    )

    doc.add_paragraph()
    doc.add_heading("Production (Planned)", level=2)
    doc.add_paragraph(
        "OAuth 2.0 Client Credentials flow will be provisioned for production use."
    )

    # Section 3: Available Endpoints
    doc.add_heading("3. Available Endpoints", level=1)
    add_table(
        doc,
        ["Endpoint", "Method", "Auth", "Description"],
        [
            ["/health", "GET", "No", "Service health check"],
            ["/api/v1/test", "GET", "Yes", "Validate authentication"],
            ["/api/v1/echo", "POST", "Yes", "Test bidirectional communication"],
        ],
    )

    doc.add_paragraph()

    # Section 4: Request & Response Examples
    doc.add_heading("4. Request & Response Examples", level=1)

    # 4.1 Health Check
    doc.add_heading("4.1 Health Check (No Auth)", level=2)
    p = doc.add_paragraph()
    p.add_run("Request:").bold = True
    doc.add_paragraph("GET https://rr.kailash.ai/health")
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Response (200 OK):").bold = True
    response1 = """{
  "status": "ok",
  "timestamp": "2025-12-04T07:44:49Z",
  "message": "Service healthy"
}"""
    doc.add_paragraph(response1)

    doc.add_paragraph()

    # 4.2 Authentication Test
    doc.add_heading("4.2 Authentication Test", level=2)
    p = doc.add_paragraph()
    p.add_run("Request:").bold = True
    req2 = """GET https://rr.kailash.ai/api/v1/test
Authorization: Basic cnJwc19jcGlfdGVzdDpLYWlsYXNoMjAyNENQSQ=="""
    doc.add_paragraph(req2)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Response (200 OK):").bold = True
    response2 = """{
  "status": "ok",
  "timestamp": "2025-12-04T07:44:50Z",
  "message": "Authentication successful for user: rrps_cpi_test"
}"""
    doc.add_paragraph(response2)

    doc.add_paragraph()

    # 4.3 Echo Test
    doc.add_heading("4.3 Echo Test (POST with Payload)", level=2)
    p = doc.add_paragraph()
    p.add_run("Request:").bold = True
    req3 = """POST https://rr.kailash.ai/api/v1/echo
Authorization: Basic cnJwc19jcGlfdGVzdDpLYWlsYXNoMjAyNENQSQ==
Content-Type: application/json

{
  "message": "Hello from Kailash",
  "test_id": "RRPS-001"
}"""
    doc.add_paragraph(req3)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Response (200 OK):").bold = True
    response3 = """{
  "status": "ok",
  "timestamp": "2025-12-04T07:44:50Z",
  "received_message": "Hello from Kailash",
  "test_id": "RRPS-001",
  "echo_hash": "127e1d79f8015388"
}"""
    doc.add_paragraph(response3)

    doc.add_paragraph()

    # Section 5: Error Handling
    doc.add_heading("5. Error Handling", level=1)
    add_table(
        doc,
        ["HTTP Code", "Meaning", "Recommended Action"],
        [
            ["200", "Success", "Process response"],
            ["401", "Unauthorized", "Verify credentials"],
            ["400", "Bad Request", "Check payload format"],
            ["500", "Server Error", "Retry with backoff"],
            ["503", "Service Unavailable", "Retry after 30s"],
        ],
    )
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Suggested Retry Policy: ").bold = True
    p.add_run("3 attempts with exponential backoff (1s, 2s, 4s)")

    doc.add_paragraph()

    # Section 6: Security Summary
    doc.add_heading("6. Security Summary", level=1)
    add_table(
        doc,
        ["Control", "Implementation"],
        [
            ["Encryption", "TLS 1.3 (CHACHA20-POLY1305, AES-256-GCM)"],
            ["Authentication", "Basic Auth (test) / OAuth 2.0 (production)"],
            ["Audit Logging", "All requests logged with timestamp and correlation ID"],
            ["Availability", "Auto-restart on failure, health monitoring enabled"],
        ],
    )

    doc.add_paragraph()

    # Section 7: Testing Checklist
    doc.add_heading("7. Testing Checklist", level=1)
    add_table(
        doc,
        ["#", "Test", "Expected Result"],
        [
            ["1", "GET /health", '200 OK, status: "ok"'],
            [
                "2",
                "GET /api/v1/test with valid credentials",
                "200 OK, authentication successful",
            ],
            ["3", "GET /api/v1/test with invalid credentials", "401 Unauthorized"],
            ["4", "POST /api/v1/echo with JSON payload", "200 OK, message echoed back"],
        ],
    )

    doc.add_paragraph()

    # Section 8: Next Steps
    doc.add_heading("8. Next Steps", level=1)
    add_table(
        doc,
        ["Party", "Action"],
        [
            ["RRPS", "Provide OAuth credentials for Integrum to connect to CPI"],
            ["Integrum", "Configure CPI connection once credentials received"],
            ["Integrum", "Provision OAuth 2.0 endpoint for production"],
        ],
    )

    doc.add_paragraph()

    # Section 9: Contact
    doc.add_heading("9. Contact", level=1)
    doc.add_paragraph(
        "For connectivity issues or questions, please contact the Integrum project team."
    )

    # Save
    output_path = "/Users/ssfoo/Documents/GitHub/rr_lead_to_cash/src/lead_to_cash/docs/RRPS_Connectivity_Documentation.docx"
    doc.save(output_path)
    print(f"Created: {output_path}")
    return output_path


if __name__ == "__main__":
    create_document()
