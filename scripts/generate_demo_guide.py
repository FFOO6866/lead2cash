"""Generate Demo Data Guide as a Word document."""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_cell_shading(cell, color_hex):
    """Set cell background color."""
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn("w:shd"), {
        qn("w:fill"): color_hex,
        qn("w:val"): "clear",
    })
    shading.append(shd)


def add_table(doc, headers, rows, col_widths=None):
    """Add a formatted table to the document."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.style = doc.styles["Normal"]
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, "2F5496")

    # Data rows
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
            if r_idx % 2 == 1:
                set_cell_shading(cell, "D6E4F0")

    # Column widths
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)

    doc.add_paragraph("")
    return table


def main():
    doc = Document()

    # Title
    title = doc.add_heading("RRPS Lead-to-Cash Demo Data Guide", level=0)
    title.runs[0].font.color.rgb = RGBColor(0x2F, 0x54, 0x96)

    doc.add_paragraph("This document describes the demo data setup for the ST Engineering end-to-end scenario, including data sources, simulated data, and environment configuration.")

    # ─── Section 1: Demo Customer ───
    doc.add_heading("1. Demo Customer", level=1)

    add_table(doc,
        ["Field", "Value"],
        [
            ["SAP Customer ID", "0022005992"],
            ["Customer Name", "ST Engineering Marine Ltd"],
            ["Sales Order", "3228005147"],
            ["IPAS Order", "1207814"],
            ["Opportunity", "Testing 1 (3000072515)"],
            ["Order Value", "EUR 62,630.00"],
            ["Engine", "1x MTU 8V2000M72 (720 KW @ 2100 RPM)"],
            ["Country", "Singapore"],
        ],
        col_widths=[5, 12],
    )

    # ─── Section 2: Data Sources ───
    doc.add_heading("2. Data Sources", level=1)

    doc.add_heading("Real Integrations", level=2)

    add_table(doc,
        ["Module", "Endpoint", "Auth Method"],
        [
            ["Credit Check", "SAP CPI — rrps-dev.it-cpi005-rt.cfapps.eu20.hana.ondemand.com\n/http/Integrum/RequestTableData", "OAuth2 client_credentials"],
            ["Opportunities", "SAP CPI — rrps-dev.it-cpi005-rt.cfapps.eu20.hana.ondemand.com\n/http/Integrum/GetOpportunity", "OAuth2 client_credentials"],
            ["Aravo KYP", "prod.aravo.co.uk/aems/restservices/v5.0/reports/25942819", "HTTP Basic Auth\n(IP-restricted to prod server)"],
        ],
        col_widths=[3.5, 10, 4],
    )

    doc.add_heading("Simulated Data", level=2)
    doc.add_paragraph("All simulated data uses real SAP document numbers sourced from the reference document (Agentic AI - Test Data.docx).")

    add_table(doc,
        ["Module", "Source File", "Data Origin"],
        [
            ["IPAS Order", "data/ipas_xml/STE_1207814.XML", "Constructed from Word doc reference data, follows IPAS SUN FORMAT v1.0 schema"],
            ["FinOps (Billing/Aging)", "src/rrps_lead2cash/services/finops_simulator.py", "SAP document numbers extracted from Word doc (BKPF/BSEG tables)"],
            ["Entity Registry", "src/rrps_lead2cash/services/entity_registry.py", "9 known customers with aliases for fuzzy name resolution"],
        ],
        col_widths=[4, 7, 6.5],
    )

    # ─── Section 3: IPAS Order ───
    doc.add_heading("3. IPAS Order Detail", level=1)
    doc.add_paragraph("File: data/ipas_xml/STE_1207814.XML")

    add_table(doc,
        ["Field", "Value"],
        [
            ["IPAS Order Number", "1207814"],
            ["SAP Sales Order (O3)", "3228005147"],
            ["Customer", "0022005992 — ST Engineering Marine Ltd"],
            ["Address", "7 Benoi Crescent, Singapore 629971"],
            ["Engine", "1x 8V2000M72 (V8, 2000-series marine propulsion)"],
            ["Power", "720 KW @ 2100 RPM"],
            ["Gross Price", "EUR 62,630.00"],
            ["Incoterms", "FOB SINGAPORE"],
            ["Payment Terms", "Z080"],
            ["Classification Society", "LR (Lloyd's Register)"],
            ["Distribution Channel", "10 (Marine)"],
            ["Delivery Date", "2025-06-30 (completed)"],
            ["Billing Plan", "Active"],
        ],
        col_widths=[5, 12],
    )

    doc.add_heading("BOM Items", level=2)

    add_table(doc,
        ["Item", "Material Number", "Description", "Qty"],
        [
            ["0001", "XM820010.00012", "MTU 8V2000M72 Marine Propulsion Engine", "1"],
            ["0002", "XM824000.00038/S", "Marine Gearbox Adapter Kit", "1"],
            ["0003", "XM830200.00007", "Marine Engine Monitoring System", "1"],
        ],
        col_widths=[2, 4.5, 8, 2],
    )

    # ─── Section 4: FinOps ───
    doc.add_heading("4. FinOps Data (Billing & Payments)", level=1)
    doc.add_paragraph("All document numbers below are real SAP references from the test data document.")

    doc.add_heading("Down Payments", level=2)

    add_table(doc,
        ["", "1st Down Payment (20%)", "2nd Down Payment (80%)"],
        [
            ["Billing Document", "3851802713", "3851802753"],
            ["Amount (EUR)", "12,526.00", "50,104.00"],
            ["Financial Document", "5130025522", "5130025815"],
            ["Clearing Document", "5140059754", "5140060121"],
            ["Clearing Date", "2025-06-20", "2025-07-30"],
            ["Status", "CLEARED", "CLEARED"],
            ["Company Code", "0011", "0011"],
        ],
        col_widths=[4.5, 6, 6],
    )

    doc.add_heading("Financial Summary", level=2)

    add_table(doc,
        ["Metric", "Value"],
        [
            ["Total Billed", "EUR 62,630.00"],
            ["Total Collected", "EUR 62,630.00"],
            ["Outstanding", "EUR 0.00"],
            ["Billing Status", "ON_TRACK"],
            ["Aging Risk", "LOW (all buckets zero)"],
        ],
        col_widths=[5, 12],
    )

    # ─── Section 5: Other Customers ───
    doc.add_heading("5. Other Demo Customers", level=1)
    doc.add_paragraph("Additional customers in the FinOps simulator for demonstrating different scenarios.")

    add_table(doc,
        ["Customer", "SAP ID", "Sales Order", "Value (EUR)", "Status", "Demo Purpose"],
        [
            ["CLLS Power System Ltd", "0021000090", "3228006201", "45,000", "ON_TRACK\n(31,500 outstanding)", "Partial payment scenario"],
            ["Tianjin Dingsheng", "0022005601", "3228007450", "5,200,000", "OVERDUE\n(100% outstanding, 90+ bucket)", "High-risk / blocked scenario"],
        ],
        col_widths=[3.5, 2.5, 2.5, 2.5, 3.5, 3],
    )

    # ─── Section 6: Demo Flow ───
    doc.add_heading("6. End-to-End Demo Flow", level=1)
    doc.add_paragraph('When a user asks "What is the status of ST Engineering?", the system executes:')

    add_table(doc,
        ["Step", "Module", "Source", "What It Returns"],
        [
            ["1", "Entity Resolution", "SIMULATED", "Resolves 'ST Engineering' to SAP ID 0022005992"],
            ["2", "Credit Check", "REAL (SAP CPI)", "Credit limit, exposure, customer address from SAP"],
            ["3", "Opportunities", "REAL (SAP CPI)", "CEC opportunity 'Testing 1' (3000072515)"],
            ["4", "Aravo KYP", "REAL (Aravo API)", "Third-party compliance status and risk rating"],
            ["5", "IPAS Order", "SIMULATED (XML)", "1x 8V2000M72 engine, EUR 62,630, 3 BOM items"],
            ["6", "FinOps Billing", "SIMULATED", "2 down payments, both CLEARED, fully paid"],
            ["7", "FinOps Aging", "SIMULATED", "Zero receivables, risk LOW"],
        ],
        col_widths=[1.5, 3, 3.5, 9.5],
    )

    # ─── Section 7: Environment ───
    doc.add_heading("7. Environment Variables", level=1)

    add_table(doc,
        ["Variable", "Purpose", "Where Set"],
        [
            ["SAP_CPI_BASE_URL", "CPI runtime base URL", "GitHub Secrets"],
            ["SAP_CPI_TOKEN_URL", "OAuth2 token endpoint", "GitHub Secrets"],
            ["SAP_CPI_CLIENT_ID", "OAuth2 client ID", "GitHub Secrets"],
            ["SAP_CPI_CLIENT_SECRET", "OAuth2 client secret", "GitHub Secrets (MISSING)"],
            ["ARAVO_REPORT_ID", "Aravo report ID (25942819)", "GitHub Secrets"],
            ["ARAVO_AUTH_TOKEN", "Base64-encoded Basic auth token", "GitHub Secrets"],
            ["ARAVO_VERIFY_SSL", "SSL verification (false for Windows)", "GitHub Secrets"],
            ["IPAS_XML_DIR", "Path to IPAS XML directory", ".env / deployment config"],
        ],
        col_widths=[4.5, 6.5, 4.5],
    )

    # ─── Section 8: Key Files ───
    doc.add_heading("8. Key Files", level=1)

    add_table(doc,
        ["File", "Purpose"],
        [
            ["data/ipas_xml/STE_1207814.XML", "ST Engineering IPAS order (mock)"],
            ["data/ipas_xml/SSZ sample.XML", "SSZ Suzhou sample order (mock)"],
            ["src/rrps_lead2cash/services/finops_simulator.py", "Billing and aging simulator"],
            ["src/rrps_lead2cash/services/ipas_xml_parser.py", "IPAS XML parser"],
            ["src/rrps_lead2cash/services/sap_cpi_client.py", "SAP CPI client (real integration)"],
            ["src/rrps_lead2cash/services/aravo_kyp_client.py", "Aravo KYP client (real integration)"],
            ["src/rrps_lead2cash/services/entity_registry.py", "Customer name-to-ID resolution"],
            ["src/rrps_lead2cash/core/gateway.py", "FastAPI routes (22 endpoints)"],
            ["src/rrps_lead2cash/core/models.py", "All Pydantic data models"],
        ],
        col_widths=[8.5, 9],
    )

    # Save
    output = "docs/RRPS Lead-to-Cash Demo Data Guide.docx"
    doc.save(output)
    print(f"Saved: {output}")


if __name__ == "__main__":
    main()
